# -*- coding: utf-8 -*-
"""结构抽取 / 大纲组装的回归测试。

这些是**纯确定性**逻辑（不联网、不调用模型），也正是整条链路里最容易
悄悄坏掉的一层：结构抽错了不会有任何报错，只会让大纲变差、让生成的 PPT
每页变成空壳。所以用内存里合成的夹具把它们钉住。

跑法::

    .venv\\Scripts\\python.exe -m unittest discover -s tests -v
"""
from __future__ import annotations
import io
import os
import sys
import unittest
from unittest import mock

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, 'src'))

from pptgen import config, llm, parse, pipeline, structure     # noqa: E402


# ══════════════════════════════════════════════════════════════
# 夹具：在内存里合成，不提交二进制文件
# ══════════════════════════════════════════════════════════════
def _h(text, level=1, slide=None, **kw):
    """一个 heading 块。`slide` 是页号（pptx 与 pdf 共用同一套前置页机制）。"""
    b = dict(type='heading', level=level, text=text)
    if slide is not None:
        b['slide'] = slide
    b.update(kw)
    return b


def _p(text, slide=None):
    b = dict(type='para', text=text)
    if slide is not None:
        b['slide'] = slide
    return b

def _tb(slide, text, size_pt, top_in, left_in=0.5, width_in=11.0):
    from pptx.util import Inches, Pt
    box = slide.shapes.add_textbox(Inches(left_in), Inches(top_in),
                                   Inches(width_in), Inches(0.9))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size_pt)
    return box


def make_pptx_bytes() -> bytes:
    """仿「Dify 介绍与实战」的版面：封面 / 目录 / 分隔页 / 正文页 / 封底。

    关键特征一个不少：分隔页有 120pt 巨号序号，正文页左上角有章节 kicker，
    正文页里有一处**比真标题字号还大**的数字噪声（30pt 的 `60,000+`）。
    """
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    # 封面
    s = prs.slides.add_slide(blank)
    _tb(s, 'DIFY · 产品介绍与实战', 13, 1.4)
    _tb(s, '演示夹具标题', 52, 1.8)
    _tb(s, '50,186', 28, 4.8)

    # 目录页
    s = prs.slides.add_slide(blank)
    _tb(s, 'CONTENTS', 18, 0.4)
    _tb(s, '目录', 29, 0.7)
    _tb(s, '第一章内容与第二章内容', 18, 2.2)

    # 两章：分隔页 + 正文页
    for num, name in (('01', '第一章'), ('02', '第二章')):
        s = prs.slides.add_slide(blank)
        _tb(s, num, 120, 1.9)
        _tb(s, 'SECTION %s' % num, 13, 2.5)
        _tb(s, name, 40, 2.9)
        _tb(s, '%s 的要点说明' % name, 15, 3.9)
        for j in range(2):
            s = prs.slides.add_slide(blank)
            _tb(s, '%s %s' % (num, name), 12, 0.4)          # kicker
            _tb(s, '%s 第 %d 节的内容主张' % (name, j + 1), 29, 0.7)
            if j == 0:
                # 比标题字号更大的数字噪声：标题启发式必须不选它
                _tb(s, '60,000+', 30, 3.0)
            _tb(s, '这一节的正文说明文字，用来撑出一点内容长度。' * 3, 15, 3.6)

    # 封底
    s = prs.slides.add_slide(blank)
    _tb(s, '让我们一起探索', 40, 2.9)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_docx_bytes() -> bytes:
    import docx
    d = docx.Document()
    d.add_heading('第一章 产品概述', level=1)
    d.add_paragraph('这是一段正文，用来说明产品的定位与目标用户。' * 3)
    d.add_heading('核心能力', level=2)
    d.add_paragraph('支持多种文档格式导入。')
    d.add_heading('第二章 技术架构', level=1)
    d.add_paragraph('解析层把文档拆成统一的内容块。' * 3)
    d.add_heading('模块划分', level=2)
    d.add_paragraph('解析、大纲、规划、渲染四段。')
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def pptx_doc() -> dict:
    return parse.parse_bytes(make_pptx_bytes(), '夹具.pptx')


def _slide_text(spec, skip=('layout', 'title', 'kicker', 'page')) -> str:
    """把一个 slide spec 里所有渲染出来的文字拼起来（用于断言「取到内容了」）。

    递归到底。早先只处理 str/list，条目里的 `{"name":…, "desc":…}` 直接掉在
    地上 —— 于是 `tinted_bands`、`numbered_columns` 这些版式在断言里等于
    「空页」，而它们内容明明都在。
    """
    out = []
    if isinstance(spec, dict):
        for k, v in spec.items():
            if k in skip:
                continue
            out.append(_slide_text(v, skip))
    elif isinstance(spec, (list, tuple)):
        out.extend(_slide_text(x, skip) for x in spec)
    elif isinstance(spec, str):
        out.append(spec)
    return ' '.join(t for t in out if t)


# ══════════════════════════════════════════════════════════════
class TestTitleHeuristic(unittest.TestCase):
    """「什么算标题」的判据。大字号 ≠ 标题。"""

    def test_rejects_number_noise(self):
        for t in ('60,000+', '01', 'IV', '3 个', '1', '50,186', '2024-10-31'):
            self.assertFalse(structure.is_title_like(t), '%r 不该算标题' % t)

    def test_accepts_real_titles(self):
        for t in ('什么是 Dify', '九大核心理念', '工作流总览：8 个节点跑通内容生产线'):
            self.assertTrue(structure.is_title_like(t), '%r 应该算标题' % t)

    def test_pptx_does_not_pick_bigger_number_as_title(self):
        """30pt 的 `60,000+` 不能压过 29pt 的真标题（实测踩过这个坑）。"""
        doc = pptx_doc()
        heads = [b['text'] for b in doc['blocks'] if b['type'] == 'heading']
        self.assertIn('第一章 第 1 节的内容主张', heads)
        self.assertNotIn('60,000+', heads)


class TestPptxStructure(unittest.TestCase):

    def setUp(self):
        self.doc = pptx_doc()
        self.sk = self.doc['structure']

    def test_chapters_from_dividers(self):
        self.assertEqual(self.sk['method'], 'divider')
        self.assertEqual([c['name'] for c in self.sk['chapters']],
                         ['01 第一章', '02 第二章'])
        self.assertEqual([len(c['pages']) for c in self.sk['chapters']], [2, 2])

    def test_front_matter_covers_whole_slides(self):
        """front matter 要标整页：封面的装饰文字不能漏进正文。"""
        blocks = self.doc['blocks']
        front_text = ' '.join(structure.block_text(blocks[i])
                              for i in self.sk['front_matter'])
        self.assertIn('50,186', front_text)          # 封面的大数字被排除了
        self.assertNotIn('50,186', parse.outline_source(self.doc))

    def test_divider_is_not_a_content_page(self):
        """分隔页是章名，不该被算成那一章的第一页。"""
        names = [p['name'] for c in self.sk['chapters'] for p in c['pages']]
        self.assertNotIn('01 第一章', names)
        self.assertNotIn('02 第二章', names)

    def test_chapter_without_divider_still_detected(self):
        """章节号集合 = 分隔页 ∪ kicker。

        实测源 deck 的第 5 章没有分隔页、只有 kicker —— 只认分隔页会少一章。
        """
        doc = self.doc
        # 造一个只靠 kicker 出现的第 3 章
        import copy
        blocks = copy.deepcopy(doc['blocks'])
        blocks.append(dict(type='heading', level=2, text='收尾', slide=99,
                           role='body', kicker='03 总结'))
        blocks.append(dict(type='para', text='落地要点与下一步。', slide=99))
        sk = structure.build_skeleton(blocks)
        self.assertEqual(len(sk['chapters']), 3)
        self.assertEqual(sk['chapters'][2]['name'], '03 总结')


class TestDocxStructure(unittest.TestCase):

    def test_depth_fallback(self):
        doc = parse.parse_bytes(make_docx_bytes(), '夹具.docx')
        sk = doc['structure']
        self.assertEqual(sk['method'], 'depth')
        self.assertEqual([c['name'] for c in sk['chapters']],
                         ['第一章 产品概述', '第二章 技术架构'])


class TestCompression(unittest.TestCase):
    """超页数时**保章压页**：章节一个不少。"""

    def test_allocate_guarantees_one_page_per_chapter(self):
        # 5 章 18 页压到 15：每章至少 1 页
        self.assertEqual(sum(pipeline._allocate([3, 3, 4, 7, 1], 15)), 15)
        self.assertTrue(all(q >= 1 for q in pipeline._allocate([3, 3, 4, 7, 1], 15)))

    def test_allocate_never_drops_chapters(self):
        for budget in range(1, 12):
            quota = pipeline._allocate([3, 3, 4, 7, 1], budget)
            self.assertTrue(all(q >= 1 for q in quota[:min(budget, 5)]))

    def test_compress_sections_keeps_every_chapter(self):
        sections = [dict(name='第%d章' % i, pages=list(range(n)))
                    for i, n in enumerate([3, 3, 4, 7, 1])]
        out = pipeline._compress_sections(sections, 15)
        self.assertEqual(len(out), 5)                       # 一章都没少
        self.assertEqual(sum(len(s['pages']) for s in out), 15)

    def test_pick_evenly_keeps_first_and_last(self):
        pages = list(range(10))
        picked = pipeline._pick_evenly(pages, 3)
        self.assertEqual(picked[0], 0)
        self.assertEqual(picked[-1], 9)

    def test_long_doc_keeps_structure(self):
        """长文档走压缩路径时，每一章都要还在。"""
        doc = pptx_doc()
        txt = parse.outline_source(doc, max_chars=800)
        for c in doc['structure']['chapters']:
            self.assertIn(c['name'], txt)


class TestPageQuota(unittest.TestCase):
    """各章正文页数**按原文分量分配**，不是每章一页。

    实测的病：12 页上限 / 6 章时，分隔页把正文预算压成 6–6 页，于是每章恰好
    一页 —— 无论上传什么文档，大纲都是「六部分、每部分两页」的同一副骨架。
    """

    @staticmethod
    def _unit(i):
        return dict(name='源页 %d' % i, start=i, end=i + 1, lead='', chars=100)

    def setUp(self):
        self.chapters = [dict(name='01 开篇', pages=[self._unit(1)]),
                         dict(name='02 主体', pages=[self._unit(i) for i in range(2, 8)]),
                         dict(name='03 收尾', pages=[self._unit(9)])]

    def _quota(self, hi=12):
        return pipeline._allocate([len(c['pages']) for c in self.chapters], hi)

    def test_quota_follows_the_source_weight(self):
        quota = self._quota()
        self.assertEqual(sum(quota), 12)
        self.assertGreater(quota[1], quota[0], '内容厚的章该多拿页')
        self.assertGreater(quota[1], quota[2])

    def test_prompt_lists_each_chapter_quota(self):
        block = pipeline._quota_block(self.chapters, self._quota())
        for c in self.chapters:
            self.assertIn(c['name'], block)
        self.assertIn('不必均等', block)

    def test_quota_block_is_empty_without_a_skeleton(self):
        self.assertEqual('', pipeline._quota_block([], []))
        self.assertEqual('', pipeline._quota_block(self.chapters, [1]))

    def test_oneshot_prompt_carries_the_quota(self):
        """短文档（一次调用）也得拿到页数分配信号，否则模型只会均分。"""
        sk = dict(chapters=self.chapters)
        doc = dict(blocks=[], source='夹具.md', title='夹具', structure=sk)
        seen = {}

        def fake(prompt, *a, **kw):
            seen['prompt'] = prompt
            return {}

        with mock.patch.object(llm, 'ask_json', side_effect=fake):
            pipeline._outline_oneshot(doc, '正文', 9, 12, object(), sk)
        self.assertIn('01 开篇：1 页', seen['prompt'])
        self.assertIn('02 主体：', seen['prompt'])


class TestOutlineAssembly(unittest.TestCase):

    def test_toc_line_is_capped(self):
        long_name = '04 基于 Dify 的开发实战'
        line = pipeline._toc_line(long_name, '一段很长很长很长很长的章节概括说明文字')
        self.assertLessEqual(len(line), pipeline._TOC_MAX)

    def test_normalise_plan_backfills_header(self):
        """规划结果缺 title/kicker 时要从大纲补齐。

        实测所有 LLM 生成的 deck 都是 has_title=0/N —— 渲染出来正文页顶部
        一片空白（`tokens.header()` 在两者都空时什么都不画）。
        """
        outline = dict(title='T', toc=[], sections=[dict(name='01 甲', summary='', pages=[
            dict(title='页一', hint=''),
            dict(title='页二', hint=''),
        ])])
        plan = pipeline._normalise_plan(
            [dict(layout='numbered_columns', items=[dict(name='a', desc='b')]),
             dict(layout='data_table', header=['h'], rows=[['r']])],
            outline)
        self.assertEqual(plan['slides'][0]['title'], '页一')
        self.assertEqual(plan['slides'][0]['kicker'], '01 甲')
        self.assertEqual(plan['slides'][1]['title'], '页二')

    def test_normalise_plan_带回副标题(self):
        """封面副标题不在大纲生成提示词里，但**按页修订会写它**。

        `revise.commit_page` 改封面副标题时同时写 deck 与 `outline['subtitle']`
        （还是为了不被下一次重新规划冲掉）。这里不带出来的话，用户改过的副标题
        会在下一次「生成 PPT」时静默变回日期兜底值 —— 别的字段都改了，只有它变回去。
        """
        sections = [dict(name='01 甲', summary='', pages=[
            dict(title='页一', hint='', source='S1')])]
        base = dict(title='T', toc=[], sections=sections)
        slides = [dict(layout='numbered_columns', items=[])]

        plan = pipeline._normalise_plan(slides, dict(base, subtitle='2026 年 9 月'))
        self.assertEqual(plan['subtitle'], '2026 年 9 月')
        # 没有这个键时不能炸，也不能凭空编一个
        self.assertEqual(pipeline._normalise_plan(slides, base)['subtitle'], '')

    def _plan_one(self, anchor: str, title: str) -> dict:
        """用「模型自创的标题」（对不上任何 heading）规划一页。

        标题对不上时唯一能救的就是锚点 —— 这正是实测的失配场景：
        LLM 大纲给的是主张式标题，`by_title` 永远查不到。
        """
        doc = pptx_doc()
        sk = doc['structure']
        outline = dict(title='T', toc=[], sections=[dict(
            name=sk['chapters'][0]['name'], summary='', pages=[
                dict(title=title, hint='', source='', anchor=anchor)])])
        return pipeline._plan_fallback(outline, doc)['slides'][0]

    def test_plan_fallback_uses_anchor(self):
        src_page = pptx_doc()['structure']['chapters'][0]['pages'][0]['name']
        with_anchor = self._plan_one(src_page, '一个自创的主张式标题')
        # 锚点生效 → 取到了该源页的正文，而不是只有标题
        self.assertIn('这一节的正文说明文字', _slide_text(with_anchor))

    def test_plan_fallback_without_anchor_degrades(self):
        """没有锚点、标题又对不上 → 退化成只有标题的一页。

        这条是「锚点为什么必要」的对照：以前正是这样静默降级的（当年是
        statement；具体用哪套版式现在由**启用清单**决定，见 `_title_only`，
        所以这里只断言「没有取到素材、产出是标题页」）。
        """
        without = self._plan_one('', '一个自创的主张式标题')
        self.assertNotIn('这一节的正文说明文字', _slide_text(without))
        self.assertIn(without['layout'], pipeline._TITLE_ONLY_ORDER)

    def test_plan_fallback_takes_a_digest_line_anchor(self):
        """anchor 抄成**整行骨架摘要**时也取得到素材。

        实测那份白皮书 13 页的 anchor 全是这种（`12 群晖白皮书　选择性同步…
        （1380 字）`），一条都对不上 → 整份 deck 的素材是空的，而它只表现为
        「掉进确定性兜底的那几页只剩一行标题」。
        """
        doc = pptx_doc()
        name = doc['structure']['chapters'][0]['pages'][0]['name']
        line = self._digest_line(doc, name)
        got = self._plan_one(line, '一个自创的主张式标题')
        self.assertIn('这一节的正文说明文字', _slide_text(got))

    def test_align_anchors_rewrites_digest_lines_only(self):
        """归一：认得出的换成源页名，认不出的**原样留着**（多半是人手写的）。

        归一要在规划之前做 —— 规划阶段靠 anchor 取素材，也靠它重算 intent。
        """
        sk = pptx_doc()['structure']
        name = sk['chapters'][0]['pages'][0]['name']
        line = self._digest_line(pptx_doc(), name)
        sections = [dict(name='01 甲', pages=[dict(title='页一', anchor=line),
                                              dict(title='页二', anchor='对不上的一行')])]
        logs = []
        bad = pipeline._align_anchors(sections, sk, log=logs.append)
        self.assertEqual(1, bad)
        self.assertEqual(name, sections[0]['pages'][0]['anchor'])
        self.assertEqual('对不上的一行', sections[0]['pages'][1]['anchor'])
        self.assertTrue(logs, '对不上时要留下日志')

    @staticmethod
    def _digest_line(doc: dict, name: str) -> str:
        """骨架摘要里那一行（`   - 页名　首句（N 字）`）。"""
        for line in structure.skeleton_digest(doc['structure']).splitlines():
            if line.strip().startswith('- ' + name):
                return line.strip()[2:]
        raise AssertionError('摘要里没有这一页：%s' % name)


class TestShortenTitles(unittest.TestCase):
    """超长标题交给模型缩写（用户明确要求「标题太长就用模型总结来缩短」）。

    为什么必须缩写而不是只靠折行：源文档的章名是「章名　—　副题」一整串
    （`structure.skeleton_digest` 就是这么拼的），30 多字印在分隔页的 40pt
    大字上，折两行也全是字。缩写是语义判断，只能模型来做 —— 但它**不许影响
    主流程**：拿不到结果就保留原文。
    """

    def _outline(self):
        return dict(
            title='一个非常长的整份 PPT 标题超过了二十个字的上限',
            sections=[dict(
                name='01 初识 Dify　—　什么是 Dify · 设计初衷 · 九大核心理念',
                summary='Dify 是融合 BaaS 与 LLMOps 的开源 LLM 应用平台。',
                pages=[dict(title='Dify 是什么：开源 LLM 应用开发平台与它的技术栈分层'),
                       dict(title='短标题')])])

    def test_long_titles_are_listed_with_budget(self):
        long = pipeline._long_titles(self._outline())
        self.assertEqual(['title', 'section:0', 'page:0:0'], [k for k, _, _ in long])
        self.assertEqual([20, 14, 24], [b for _, _, b in long])
        self.assertEqual([], pipeline._long_titles(
            dict(title='短标题', sections=[dict(name='01 甲', pages=[dict(title='乙')])])))

    def test_model_shrinks_titles_and_keeps_the_chapter_number(self):
        ol = self._outline()
        with mock.patch.object(llm, 'ask_json', return_value={
                'titles': ['Dify 介绍与实战', '初识 Dify', 'Dify 是什么']}):
            pipeline._shorten_titles(ol, object(), log=lambda m: None)
        self.assertEqual('Dify 介绍与实战', ol['title'])
        self.assertEqual('01 初识 Dify', ol['sections'][0]['name'],
                         '章节号不能被缩写吃掉 —— 分隔页的大号编号靠它')
        self.assertEqual(['Dify 是什么', '短标题'],
                         [p['title'] for p in ol['sections'][0]['pages']])

    def test_model_output_written_as_dicts_is_accepted(self):
        ol = self._outline()
        with mock.patch.object(llm, 'ask_json', return_value={
                'titles': [{'text': 'Dify 介绍与实战'}, {'text': '01 初识 Dify'},
                           {'text': 'Dify 是什么'}]}):
            pipeline._shorten_titles(ol, object(), log=lambda m: None)
        self.assertEqual('01 初识 Dify', ol['sections'][0]['name'])

    def test_wrong_count_keeps_everything(self):
        ol = self._outline()
        before = ol['sections'][0]['name']
        with mock.patch.object(llm, 'ask_json', return_value={'titles': ['只有一个']}):
            pipeline._shorten_titles(ol, object(), log=lambda m: None)
        self.assertEqual(before, ol['sections'][0]['name'])
        self.assertEqual('一个非常长的整份 PPT 标题超过了二十个字的上限', ol['title'])

    def test_llm_failure_keeps_everything_and_does_not_raise(self):
        ol = self._outline()
        before = [ol['title'], ol['sections'][0]['name']]
        with mock.patch.object(llm, 'ask_json',
                               side_effect=llm.LLMError('超时')):
            pipeline._shorten_titles(ol, object(), log=lambda m: None)
        self.assertEqual(before, [ol['title'], ol['sections'][0]['name']])

    def test_not_actually_shorter_is_rejected(self):
        """模型原样抄回来（甚至压得更长）不算做成 —— 保留原文。"""
        ol = self._outline()
        before = ol['sections'][0]['name']
        with mock.patch.object(llm, 'ask_json', return_value={
                'titles': [ol['title'], before, '更' * 40]}):
            pipeline._shorten_titles(ol, object(), log=lambda m: None)
        self.assertEqual(before, ol['sections'][0]['name'])
        self.assertEqual('Dify 是什么：开源 LLM 应用开发平台与它的技术栈分层',
                         ol['sections'][0]['pages'][0]['title'])

    def test_no_long_title_means_no_model_call(self):
        ol = dict(title='短', sections=[dict(name='01 甲', pages=[dict(title='乙')])])
        with mock.patch.object(llm, 'ask_json') as m:
            pipeline._shorten_titles(ol, object(), log=lambda m: None)
        self.assertFalse(m.called, '没有超长标题就不该花一次模型调用')


class TestSkeletonDigestSplitsSubtitle(unittest.TestCase):
    """骨架里章名与副题必须**分两行**写。

    早先拼成一行 `## 01 初识 Dify　—　什么是 Dify · 设计初衷 · 九大核心理念`，
    模型就照抄整行当章节名 —— 那个名字要印在分隔页的 40pt 大字上，
    一行只放得下约 14 字，必然被截成 `01 初识 Dify　—　什么是 Di…`。
    """

    def test_name_and_subtitle_are_on_separate_lines(self):
        sk = dict(chapters=[dict(name='01 初识 Dify', subtitle='什么是 Dify · 设计初衷',
                                 pages=[dict(name='什么是 Dify', lead='', chars=10)])])
        txt = structure.skeleton_digest(sk)
        first = txt.splitlines()[0]
        self.assertEqual('## 01 初识 Dify', first)
        self.assertNotIn('—', first, '章名那一行不许带破折号拼出来的副题')
        self.assertIn('设计初衷', txt)


class TestOutlineSignal(unittest.TestCase):
    """文档自带的目录（PDF 书签树）优先于一切版面推断。"""

    BLOCKS = [
        _h('Dify: Product Introduction', 1, 1, role='cover'),
        _p('Dify 是一个开源的 LLM 应用开发平台。', 1),
        _h('目录', 1, 2, role='toc'),
        _p('前言 ............. 3', 2),
        _h('前言', 1, 3),
        _p('还记得 2013 年，我们创立了乐豆信息。', 3),
        _h('1 初识 Dify', 1, 4),
        _p('Dify 是一个多合一的数据处理与分析平台。', 4),
        _h('1.1 什么是 Dify', 2, 4),
        _p('本节解释 Dify 的定位。', 4),
        _h('2 为什么选 Dify', 1, 5),
        _p('AI 应用开发的四大挑战。', 5),
    ]
    OUTLINE = [dict(level=0, title='目录', page=2),
               dict(level=0, title='前言', page=3),
               dict(level=0, title='1 初识 Dify', page=4),
               dict(level=0, title='2 为什么选 Dify', page=5)]

    def setUp(self):
        self.sk = structure.build_skeleton(self.BLOCKS, 'pdf', self.OUTLINE)

    def test_method_is_outline(self):
        self.assertEqual(self.sk['method'], 'outline')

    def test_chapters_come_from_the_outline(self):
        self.assertEqual([c['name'] for c in self.sk['chapters']],
                         ['01 前言', '02 初识 Dify', '03 为什么选 Dify'])

    def test_toc_entry_is_not_a_chapter(self):
        """目录条目不是章 —— 目录页的内容是各章标题粘在一起的产物。"""
        self.assertNotIn('目录', [c['name'] for c in self.sk['chapters']])

    def test_outline_wins_over_page_roles(self):
        """书签说的是真章节，就不该再被「每页一个页单元」的粒度带跑。"""
        self.assertEqual([len(c['pages']) for c in self.sk['chapters']], [1, 2, 1])

    def test_without_outline_it_falls_back(self):
        sk = structure.build_skeleton(self.BLOCKS, 'pdf', None)
        self.assertNotEqual(sk['method'], 'outline')

    def test_placeholder_bookmark_is_not_a_chapter(self):
        """`书签 1` 是 Acrobat 给没命名的书签自动起的名字，不是章节。

        实测那份 Synology 白皮书的第一条书签就是它，落在一页没有标题的页上：
        留着它，大纲会给它编一页正文、分隔页也照插一页，白吃两页预算。
        它的标题块也不能漏下去当页单元 —— 那会变成一张标题叫「書籤 1」的页。
        """
        blocks = [_h('書籤 1', 1, 2),
                  _h('Introduction', 1, 3), _p('正文若干字。' * 20, 3),
                  _h('Architecture', 1, 4), _p('正文若干字。' * 20, 4),
                  _h('Security', 1, 5), _p('正文若干字。' * 20, 5)]
        outline = [dict(level=0, title='書籤 1', page=2),
                   dict(level=0, title='Introduction', page=3),
                   dict(level=0, title='Architecture', page=4),
                   dict(level=0, title='Security', page=5)]
        sk = structure.build_skeleton(blocks, 'pdf', outline)
        self.assertEqual([c['name'] for c in sk['chapters']],
                         ['01 Introduction', '02 Architecture', '03 Security'])
        names = [p['name'] for c in sk['chapters'] for p in c['pages']]
        self.assertNotIn('書籤 1', names)

    def test_empty_chapter_is_dropped_and_the_rest_renumbered(self):
        """两条书签落在同一页 → 头一条的区间宽度是 0，它是个空章。

        空章要丢掉，**后面的章号要重排**：否则第一页分隔页上印的是「02」，
        而整份 deck 里根本没有 01。
        """
        blocks = [_h('封面', 1, 1, role='cover'), _p('封面装饰。', 1),
                  _h('Introduction', 1, 2), _p('正文若干字。' * 20, 2),
                  _h('Architecture', 1, 3), _p('正文若干字。' * 20, 3)]
        outline = [dict(level=0, title='概述', page=2),      # 与下一条同页
                   dict(level=0, title='Introduction', page=2),
                   dict(level=0, title='Architecture', page=3)]
        sk = structure.build_skeleton(blocks, 'pdf', outline)
        self.assertEqual([c['name'] for c in sk['chapters']],
                         ['01 Introduction', '02 Architecture'])
        self.assertEqual([len(c['pages']) for c in sk['chapters']], [1, 1])


class TestSignalRanking(unittest.TestCase):
    """认得出来的信号谁压谁。"""

    def test_depth_outranks_numeric(self):
        """`1. 定义：` 这种列表项不该被数字前缀当成章（实测会多出两章）。"""
        blocks = [_h('第一章 总则', 1, 1), _h('1. 目的：', 2, 1), _p('正文。', 1),
                  _h('第二章 范围', 1, 2), _h('2. 平台：', 2, 2), _p('正文。', 2)]
        sk = structure.build_skeleton(blocks)
        self.assertEqual(sk['method'], 'depth')
        self.assertEqual([c['name'] for c in sk['chapters']], ['第一章 总则', '第二章 范围'])

    def test_front_roles_are_never_chapters(self):
        """走纯文本这条路时，封面/目录的标题也不能当章。"""
        blocks = [_h('迈胜医疗设备有限公司', 1, 1, role='cover'), _p('封面。', 1),
                  _h('第一章 总则', 1, 2), _p('正文。', 2),
                  _h('第二章 范围', 1, 3), _p('正文。', 3)]
        sk = structure.build_skeleton(blocks)
        self.assertEqual([c['name'] for c in sk['chapters']], ['第一章 总则', '第二章 范围'])

    def test_flat_drops_number_only_page_names(self):
        """兜底路径也要挡页码 —— 但不能连 `前言` 这种两字标题一起挡掉。"""
        blocks = [_h('1', 1, 1), _p('正文。', 1), _h('前言', 2, 2), _p('正文。', 2)]
        sk = structure.build_skeleton(blocks)
        self.assertEqual(sk['method'], 'flat')
        self.assertEqual([p['name'] for p in sk['chapters'][0]['pages']], ['前言'])

    def test_first_chapter_absorbs_leading_content(self):
        """封面之后、第一个章节标记之前的内容要归到第 1 章，不能掉在章节之外。"""
        blocks = [_h('封面大标题', 1, 1, role='cover'), _p('封面装饰。', 1),
                  _h('产品介绍', 2, 2), _p('这是开篇的产品介绍内容。', 2),
                  _h('01 甲', 1, 3), _p('甲的内容。', 3),
                  _h('02 乙', 1, 4), _p('乙的内容。', 4)]
        sk = structure.build_skeleton(blocks)
        first = sk['chapters'][0]
        self.assertEqual(first['start'], 2)
        self.assertIn('产品介绍', [p['name'] for p in first['pages']])


class TestChapterSegmentation(unittest.TestCase):
    """骨架分不出章（或分得太碎）时，让模型把**已有页单元**归成几章。

    为什么值得有：这是「上传什么文档都只有一个章节」的通用兜底 —— 确定性信号
    抽不出来时（纯正文的 txt、没书签也没编号的 pdf），早先整份文档就是一整章。
    """

    def _doc(self, n=20) -> dict:
        blocks = []
        for i in range(1, n + 1):
            blocks += [_h('第 %d 节 标题' % i, 2, i), _p('正文若干字。' * 5, i)]
        return dict(blocks=blocks, source='夹具.md', title='夹具',
                    structure=structure.build_skeleton(blocks))

    def _run(self, doc, answer=None, error=None):
        kw = dict(side_effect=error) if error else dict(return_value=answer)
        with mock.patch.object(llm, 'ask_json', **kw):
            pipeline._maybe_segment(doc, 15, object(), lambda m: None)
        return doc['structure']

    def test_too_many_chapters_get_regrouped(self):
        doc = self._doc(20)
        self.assertEqual(len(doc['structure']['chapters']), 20)   # 太碎：20 章
        sk = self._run(doc, {'chapters': [{'name': '开篇', 'from': 1, 'to': 5},
                                          {'name': '主体', 'from': 6, 'to': 15},
                                          {'name': '收尾', 'from': 16, 'to': 20}]})
        self.assertEqual(sk['method'], 'llm_segment')
        self.assertEqual([c['name'] for c in sk['chapters']], ['01 开篇', '02 主体', '03 收尾'])
        self.assertEqual(sk['chapters'][0]['start'], 0)
        self.assertEqual(sk['chapters'][-1]['end'], len(doc['blocks']))

    def test_no_page_unit_is_lost(self):
        """分章只许重新分组，页单元一个都不能少 —— 少了就是内容对不上锚点。"""
        doc = self._doc(20)
        before = [p['name'] for c in doc['structure']['chapters'] for p in c['pages']]
        sk = self._run(doc, {'chapters': [{'name': '甲', 'from': 1, 'to': 10},
                                          {'name': '乙', 'from': 11, 'to': 20}]})
        after = [p['name'] for c in sk['chapters'] for p in c['pages']]
        self.assertEqual(before, after)

    def test_chapter_number_written_by_model_is_stripped(self):
        doc = self._doc(8)
        sk = self._run(doc, {'chapters': [{'name': '01 开篇', 'from': 1, 'to': 4},
                                          {'name': '02 收尾', 'from': 5, 'to': 8}]})
        self.assertEqual([c['name'] for c in sk['chapters']], ['01 开篇', '02 收尾'])

    def test_bad_ranges_are_rejected_and_split_evenly(self):
        """区间有漏/有重/不连续 → 整份丢弃，改用均分（宁可不好看，不能错位）。"""
        for bad in ({'chapters': [{'name': '甲', 'from': 1, 'to': 5},
                                  {'name': '乙', 'from': 7, 'to': 20}]},      # 漏了 6
                    {'chapters': [{'name': '甲', 'from': 2, 'to': 20}]},      # 不从 1 起
                    {'chapters': [{'name': '甲', 'from': 1, 'to': 20}]},      # 只有 1 章
                    {'chapters': 'nope'},                                     # 结构不对
                    {}):                                                      # 什么都没有
            doc = self._doc(20)
            sk = self._run(doc, bad)
            self.assertEqual(sk['method'], 'even_split')
            # 均分几章由页数预算推（15 页上限 → 5 章），不是老早那个写死的 6
            self.assertEqual(len(sk['chapters']), pipeline._segment_target(15))
            self.assertEqual(sum(len(c['pages']) for c in sk['chapters']), 20)

    def test_llm_failure_still_keeps_every_page(self):
        doc = self._doc(20)
        sk = self._run(doc, error=llm.LLMError('超时'))
        self.assertEqual(sk['method'], 'even_split')
        self.assertEqual(sum(len(c['pages']) for c in sk['chapters']), 20)
        self.assertEqual(sk['chapters'][-1]['end'], len(doc['blocks']))

    def test_single_chapter_without_model_is_left_alone(self):
        """只有 1 章是**结构信号缺失**，怎么分是语义判断 —— 没模型就别硬凑章数。

        用中文两字标题（`前言`/`背景`/`方案`/`计划`）造这个场景：它们在
        `_title_ok`（≥3 字）那道坎上全被滤掉，两条纯文本路都收不齐标记，
        于是整份文档压成一章 —— 正是要兜底的那种样子。
        """
        blocks = []
        for i, t in enumerate(('前言', '背景', '方案', '计划'), 1):
            blocks += [_h(t, 1, i), _p('正文若干字。' * 20, i)]
        doc = dict(blocks=blocks, source='x.txt', title='x',
                   structure=structure.build_skeleton(blocks))
        self.assertEqual(len(doc['structure']['chapters']), 1)
        self.assertEqual(len(doc['structure']['chapters'][0]['pages']), 4)
        pipeline._maybe_segment(doc, 15, None, lambda m: None)
        self.assertEqual(doc['structure']['method'], 'flat')
        self.assertEqual(len(doc['structure']['chapters']), 1)

    def test_healthy_chapter_count_is_not_touched(self):
        doc = self._doc(20)
        doc['structure']['chapters'] = doc['structure']['chapters'][:5]
        with mock.patch.object(llm, 'ask_json') as m:
            pipeline._maybe_segment(doc, 15, object(), lambda m: None)
        self.assertFalse(m.called, '章数正常就不该多花一次模型调用')

    def test_switch_off_means_no_model_call(self):
        doc = self._doc(20)
        with mock.patch.object(config, 'outline_segment', return_value=False), \
                mock.patch.object(llm, 'ask_json') as m:
            pipeline._maybe_segment(doc, 15, object(), lambda m: None)
        self.assertFalse(m.called)
        self.assertEqual(len(doc['structure']['chapters']), 20)

    def test_segment_target_follows_the_page_budget(self):
        """目标章数按页数预算推（一章 = 一页分隔 + 两页正文）。

        早先这里是写死的 `min(cap, 6)`，而模型**每次都取上限**：不管什么文档
        都分成 6 章——散架的 104 章和规规矩矩的 15 章，出口一模一样。
        """
        self.assertEqual(4, pipeline._segment_target(12))
        self.assertEqual(5, pipeline._segment_target(15))
        self.assertEqual(6, pipeline._segment_target(18))
        self.assertEqual(2, pipeline._segment_target(4))        # 下限兜到 2 章

    def test_segmentation_prompt_asks_for_the_budget_derived_count(self):
        doc = self._doc(20)
        prompts = []

        def fake(prompt, *a, **kw):
            prompts.append(prompt)
            return {'chapters': [{'name': '开篇', 'from': 1, 'to': 10},
                                 {'name': '收尾', 'from': 11, 'to': 20}]}

        with mock.patch.object(llm, 'ask_json', side_effect=fake):
            pipeline._maybe_segment(doc, 12, object(), lambda m: None)
        self.assertTrue(prompts)
        self.assertIn('归成 3–4 章', prompts[0])          # 12 页上限 → 最多 4 章
        self.assertIn('宁少勿多', prompts[0])

    def test_segmentation_runs_before_the_divider_budget(self):
        """顺序是有讲究的：章节数决定分隔页占多少页预算。"""
        order = []
        doc = self._doc(20)

        def seg(*a, **k):
            order.append('segment')

        def budget(*a, **k):
            order.append('budget')
            return (0, 10, 15)

        outline = dict(sections=[dict(name='01 甲', summary='', pages=[
            dict(title='页一', hint='h', intent='statement', source='s', anchor='')])])
        with mock.patch.object(pipeline, '_maybe_segment', side_effect=seg), \
                mock.patch.object(pipeline, '_divider_budget', side_effect=budget), \
                mock.patch.object(pipeline, '_outline_by_llm', return_value=outline), \
                mock.patch.object(config, 'llm_config',
                                  return_value=mock.Mock(model='fixture')):
            pipeline.make_outline(doc, on_log=lambda m: None)
        self.assertEqual(order, ['segment', 'budget'])


class TestValidGroups(unittest.TestCase):
    """模型返回的分组校验（纯函数）。"""

    def test_accepts_a_covering_partition(self):
        got = pipeline._valid_groups([{'name': '开篇', 'from': 1, 'to': 4},
                                      {'name': '主体', 'from': 5, 'to': 10}], 10, 7)
        self.assertEqual([('开篇', 1, 4), ('主体', 5, 10)], got)

    def test_two_char_names_are_fine(self):
        """`开篇`、`总则` 这种两字章名很正常，别套 `is_title_like`（它要 ≥4 个字）。"""
        self.assertIsNotNone(pipeline._valid_groups(
            [{'name': '总则', 'from': 1, 'to': 5}, {'name': '附则', 'from': 6, 'to': 9}], 9, 7))

    def test_rejects_degenerate_names(self):
        for name in ('目录', '1', '', 'SECTION 1', 'x' * 40):
            self.assertIsNone(pipeline._valid_groups(
                [{'name': name, 'from': 1, 'to': 5},
                 {'name': '乙', 'from': 6, 'to': 9}], 9, 7), name)

    def test_rejects_too_many_chapters(self):
        raw = [{'name': '章%s' % i, 'from': i, 'to': i} for i in range(1, 11)]
        self.assertIsNone(pipeline._valid_groups(raw, 10, 7))

    def test_even_groups_covers_everything(self):
        for n, k in ((20, 7), (5, 2), (9, 9), (3, 7)):
            groups = pipeline._even_groups(n, k)
            self.assertEqual(groups[0][0], 1)
            self.assertEqual(groups[-1][1], n)
            for a, b in groups:
                self.assertLessEqual(a, b)
            for (_, prev), (nxt, _) in zip(groups, groups[1:]):
                self.assertEqual(nxt, prev + 1, '均分也要首尾相接')


if __name__ == '__main__':
    unittest.main()

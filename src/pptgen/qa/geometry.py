# -*- coding: utf-8 -*-
"""几何 QA —— 纯程序化、不需要渲染。

为什么必须有这个：`officecli view issues` 只检查形状是否越过**右边界**，
以下这些都查不出来（实测）：
  - 文字溢出文本框
  - 两个文本框重叠
  - 字号低于下限
  - 形状远远越过**下边界**（曾出现发丝线高 9525 英寸的 bug，officecli 报 0 issues）

单位换算：EMU → 英寸 = /914400，EMU → pt = /12700
"""
from __future__ import annotations
import math

from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400.0
EMU_PT = 12700.0

# 版心右边界（避让模板装饰弧线）
SAFE_RIGHT = 12.00
MIN_FONT_PT = 12.0

# 估算参数
CJK_WIDTH = 1.00     # 汉字宽度 ≈ 字号 × 1.0
LATIN_WIDTH = 0.52   # 拉丁/数字 ≈ 字号 × 0.52
LINE_HEIGHT = 1.42   # 行高倍数（默认单倍行距下的近似值）
OVERFLOW_TOL = 1.06  # 估算高度超过可用高度 6% 以上才报
EMPTY_WARN = 0.42    # 内容带内最大连续空白块占比超过此值 → 提示"留白偏大"


def _iter_shapes(shapes):
    """递归展开 GROUP（shape_type == 6）。"""
    for sh in shapes:
        if sh.shape_type == 6:  # GROUP
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def _text(sh) -> str:
    try:
        return sh.text_frame.text if sh.has_text_frame else ''
    except Exception:
        return ''


def _bbox(sh):
    return (sh.left / EMU_IN, sh.top / EMU_IN,
            (sh.left + sh.width) / EMU_IN, (sh.top + sh.height) / EMU_IN)


def _est_text_width_pt(text: str, size_pt: float) -> float:
    w = 0.0
    for ch in text:
        w += size_pt * (CJK_WIDTH if ord(ch) > 0x2E80 else LATIN_WIDTH)
    return w


def _paragraph_widths_pt(sh):
    """按 **run 各自的字号** 估算每个段落的宽度。

    早先的版本用段内最大字号算整串宽度，会把「72pt 的数字 + 22pt 的单位」
    这样的混合段算宽一倍，产生 text_overflow 误报（s4 的大数字页就是被这样误报的）。
    """
    widths = []
    try:
        for p in sh.text_frame.paragraphs:
            w = 0.0
            for r in p.runs:
                sz = r.font.size.pt if r.font.size is not None else 18.0
                w += _est_text_width_pt(r.text, sz)
            if w:
                widths.append((w, max((r.font.size.pt for r in p.runs
                                       if r.font.size is not None), default=18.0)))
    except Exception:
        pass
    return widths


def _run_sizes(sh):
    sizes = []
    try:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.font.size is not None:
                    sizes.append(r.font.size.pt)
    except Exception:
        pass
    return sizes


def analyse(path: str, skip: set[int] | None = None) -> dict:
    """skip: 跳过的页码集合。

    默认跳过公司模板自带的封面 / 目录 / 封底 —— 它们的字号与占位符几何由模板决定
    （例如封面那行 10pt 的 "Template DID: ..."），不属于本项目的生成范围，
    报出来只会淹没有效信号。
    """
    prs = Presentation(path)
    n = len(prs.slides)
    if skip is None:
        skip = {1, 2, n} if n >= 4 else set()
    sw, sh_ = prs.slide_width / EMU_IN, prs.slide_height / EMU_IN
    issues: list[dict] = []
    per_slide = []

    for idx, slide in enumerate(prs.slides, 1):
        if idx in skip:
            continue
        shapes = list(_iter_shapes(slide.shapes))
        boxes = []
        content_boxes = []
        for sh in shapes:
            try:
                x0, y0, x1, y1 = _bbox(sh)
            except Exception:
                continue
            txt = _text(sh).strip()
            sizes = _run_sizes(sh)
            w_in, h_in = x1 - x0, y1 - y0
            name = sh.name

            def add(sev, kind, detail):
                issues.append(dict(severity=sev, slide=idx, shape=name,
                                   kind=kind, detail=detail))

            # ① 越出画布
            if x0 < -0.01 or y0 < -0.01 or x1 > sw + 0.01 or y1 > sh_ + 0.01:
                add('error', 'out_of_canvas',
                    '包围盒 (%.2f,%.2f)-(%.2f,%.2f) 越出画布 %.2fx%.2f'
                    % (x0, y0, x1, y1, sw, sh_))
            # ② 尺寸荒诞（曾在发丝线上踩到：高度 9525 英寸）
            elif w_in > sw * 1.2 or h_in > sh_ * 1.2:
                add('error', 'absurd_size',
                    '尺寸异常 %.2f x %.2f 英寸，疑似单位换算错误' % (w_in, h_in))
            # ③ 侵入弧线区（右侧安全区）
            elif x1 > SAFE_RIGHT + 0.01 and w_in < sw * 0.9:
                add('warn', 'past_safe_area',
                    '右边界 %.2f 越过安全区 %.2f（装饰弧线区）' % (x1, SAFE_RIGHT))
            # ④ 字号下限
            for sz in sizes:
                if sz < MIN_FONT_PT - 0.01:
                    add('error', 'font_below_floor',
                        '字号 %.1fpt 低于下限 %.0fpt' % (sz, MIN_FONT_PT))
                    break
            # ⑤ 文字溢出估算（按 run 各自字号累加宽度）
            pw = _paragraph_widths_pt(sh)
            if txt and pw and h_in > 0.01 and w_in > 0.01:
                avail_w_pt = w_in * 72.0
                lines, need_pt = 0.0, 0.0
                for w_pt, size in pw:
                    n_ln = max(1, math.ceil(w_pt / max(avail_w_pt, 1.0)))
                    lines += n_ln
                    need_pt += n_ln * size * LINE_HEIGHT
                if need_pt > h_in * 72.0 * OVERFLOW_TOL:
                    add('warn', 'text_overflow',
                        '估算需 %.0f 行 / %.0fpt，框高仅 %.0fpt'
                        % (lines, need_pt, h_in * 72.0))
            # ⑥ 文本框内边距会破坏对齐（zcode 明确要求 margin=0）
            if txt:
                try:
                    tf = sh.text_frame
                    if any(v and v > 0 for v in (tf.margin_left, tf.margin_right)):
                        add('warn', 'textbox_margin',
                            '文本框左右内边距非 0，会与形状/线条对不齐')
                except Exception:
                    pass

            # 表格 / 图表是 GraphicFrame，没有 text_frame，但确实是页面内容。
            # 早先只统计含文字的形状，导致三个表格页被误判为「75% 空白」。
            is_graphic = False
            try:
                is_graphic = bool(getattr(sh, 'has_table', False)
                                  or getattr(sh, 'has_chart', False))
            except Exception:
                pass

            if txt:
                boxes.append((name, x0, y0, x1, y1))
            if txt or is_graphic:
                content_boxes.append((name, x0, y0, x1, y1))

        # ⑦ 含文字形状两两重叠
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                n1, ax0, ay0, ax1, ay1 = boxes[i]
                n2, bx0, by0, bx1, by1 = boxes[j]
                ox = min(ax1, bx1) - max(ax0, bx0)
                oy = min(ay1, by1) - max(ay0, by0)
                if ox > 0.05 and oy > 0.05:
                    issues.append(dict(
                        severity='warn', slide=idx, shape='%s × %s' % (n1, n2),
                        kind='text_overlap',
                        detail='文字块重叠 %.2f x %.2f 英寸' % (ox, oy)))

        # ⑧ 留白启发式（"生硬"的可执行代理指标）
        empty_ratio = _largest_empty_block(content_boxes, sh_)
        per_slide.append(dict(slide=idx, shapes=len(shapes),
                              text_shapes=len(boxes),
                              empty_ratio=round(empty_ratio, 3)))
        if empty_ratio > EMPTY_WARN:
            issues.append(dict(
                severity='warn', slide=idx, shape='-', kind='large_empty_area',
                detail='内容带内最大连续空白块约占 %.0f%%，观感可能偏空'
                       % (empty_ratio * 100)))

    errors = sum(1 for i in issues if i['severity'] == 'error')
    warns = sum(1 for i in issues if i['severity'] == 'warn')
    return dict(path=path, slides=len(prs.slides),
                canvas=[sw, sh_], issues=issues,
                per_slide=per_slide,
                summary=dict(error=errors, warn=warns, total=len(issues)))


def _largest_empty_block(boxes, slide_h, cols=6, rows=4,
                         band=(2.00, 6.95)):
    """把内容带切成 cols×rows 网格，返回最大连续空白块占内容带的比例。"""
    y0, y1 = band
    grid = [[False] * cols for _ in range(rows)]
    for _, bx0, by0, bx1, by1 in boxes:
        for r in range(rows):
            cy0 = y0 + (y1 - y0) * r / rows
            cy1 = y0 + (y1 - y0) * (r + 1) / rows
            for c in range(cols):
                cx0 = 0.67 + (12.00 - 0.67) * c / cols
                cx1 = 0.67 + (12.00 - 0.67) * (c + 1) / cols
                if min(bx1, cx1) - max(bx0, cx0) > 0.1 and min(by1, cy1) - max(by0, cy0) > 0.1:
                    grid[r][c] = True
    seen = [[False] * cols for _ in range(rows)]
    best = 0
    for r in range(rows):
        for c in range(cols):
            if grid[r][c] or seen[r][c]:
                continue
            stack, n = [(r, c)], 0
            seen[r][c] = True
            while stack:
                cr, cc = stack.pop()
                n += 1
                for dr, dc in ((1, 0), (-1, 0), (0, 1), (0, -1)):
                    nr, nc = cr + dr, cc + dc
                    if 0 <= nr < rows and 0 <= nc < cols \
                            and not grid[nr][nc] and not seen[nr][nc]:
                        seen[nr][nc] = True
                        stack.append((nr, nc))
            best = max(best, n)
    return best / float(cols * rows)


def format_report(rep: dict) -> str:
    lines = ['几何 QA — %s' % rep['path'],
             '页数 %d ｜ 画布 %.2f x %.2f ｜ 安全右边界 %.2f'
             % (rep['slides'], rep['canvas'][0], rep['canvas'][1], SAFE_RIGHT),
             '-' * 68]
    if rep['summary']['total'] == 0:
        lines.append('未发现问题 [OK]')
    for it in rep['issues']:
        mark = 'ERR ' if it['severity'] == 'error' else 'warn'
        lines.append('[%s] s%-2d %-22s %-18s %s'
                     % (mark, it['slide'], it['shape'][:22],
                        it['kind'], it['detail']))
    s = rep['summary']
    lines.append('-' * 68)
    lines.append('合计: %d error, %d warn' % (s['error'], s['warn']))
    return '\n'.join(lines)

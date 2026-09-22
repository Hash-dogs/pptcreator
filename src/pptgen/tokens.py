# -*- coding: utf-8 -*-
"""设计令牌与绘图原语。

本文件是全部版式的单一真相来源。所有尺寸单位为**英寸**。

三条硬约束（来自最小验证的实测结论，不要违反）：

1. **只用 `Blank` / `Title Only` layout。**
   `Full Blank` 上有个 `Rectangle 3`（x=9.19, y=0, 4.15x1.24，纯色填充）会盖住母版的
   MEVION 徽标，实测渲染出来是纯白页 —— 没有任何公司品牌元素。

2. **内容右边界 12.00"。**
   装饰弧线 `Picture` 位于 x=10.19~15.13（图片框伸出画布外），实测墨迹自 x≈12.00 起。

3. **每个 run 必须显式设置字体**（`font.name` + `<a:ea>`）。
   公司模板的主题 `fontScheme` 里 `ea` 为空，不显式设置中文会走系统默认。
"""
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ── 画布与版心 ────────────────────────────────────────────────
CANVAS_W, CANVAS_H = 13.33, 7.50   # 公司模板实际为 13.333 x 7.5
LEFT, RIGHT = 0.67, 12.00          # 版心左右边界
W = RIGHT - LEFT                   # 11.33

# 纵向分区
Y_KICKER = 0.38
Y_TITLE = 0.68
Y_CONTENT = 2.00
Y_CONTENT_BOTTOM = 6.95
Y_SOURCE = 6.78
Y_PAGENUM = 7.18

# ── 品牌色（取自模板 theme1.xml 的 clrScheme）─────────────────
RED   = RGBColor(0xD3, 0x12, 0x45)   # accent1 —— 强调，只打在单一焦点上
# accent4。**版式里不要再拿它当结构色** —— 模板自己的页面从不用蓝（theme1 里
# accent4 只出现在一个空段落的 endParaRPr 上，是残留不是设计）。早先表格表头、
# 节点框、支撑数字都用了它，成品里那几页明显跳出色系。留着这份定义只是为了让
# `resolve_color('BLUE')` 仍然可用，别在 layouts.py 里引用。
BLUE  = RGBColor(0x00, 0x46, 0x7F)   # accent4 —— 未使用
DARK  = RGBColor(0x23, 0x1F, 0x20)   # accent3 —— 正文
MUTED = RGBColor(0x64, 0x64, 0x63)   # accent5 —— 次要
GREY  = RGBColor(0x80, 0x7F, 0x83)   # accent2
RULE  = RGBColor(0xC9, 0xC9, 0xC9)   # 发丝线
TINT  = RGBColor(0xEE, 0xF3, 0xF8)   # 浅色带
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# ── 字体与字号 ────────────────────────────────────────────────
EA, LAT = '微软雅黑', 'Segoe UI'
# 模板自己的展示字体：封面标题（Helvetica Light 那一档）与目录页标题都用它。
# 品牌页保持品牌的字重，正文页才用上面那对。
EA_LIGHT = '微软雅黑 Light'
FS = dict(kicker=12, title=26, h2=18, h3=17, body=16, small=13, source=12, min=12)

# 公司骨架页在模板里的 layout 名（封面/目录页不靠 layout 名定位，靠占位符 idx）
LAYOUT_BLANK = 'Blank'


# ══════════════════════════════════════════════════════════════
# 绘图原语
# ══════════════════════════════════════════════════════════════

def _rect(slide, x, y, w, h, color, *, outline=False, radius=None):
    """基础矩形。x/y/w/h 全部是**英寸**。

    ⚠️ 曾踩过的坑：早期版本对四个尺寸统一套 `Inches()`，而调用方传的是 `Pt(0.75)`
       （已是 EMU 值），于是 `Inches(9525)` 得到 9525 英寸高的巨型矩形 —— 所有发丝线
       都变成覆盖整页的色块。**officecli 的 view issues 查不出这个**（它只检查是否
       越过右边界），只有自写的边界扫描能抓到。所以：这里只接受英寸，转换在调用方做。
    """
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if outline:
        sh.fill.background()
        sh.line.color.rgb = color
        sh.line.width = Pt(0.9)
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = color
        sh.line.fill.background()
    sh.shadow.inherit = False
    if radius:
        sh.adjustments[0] = radius
    return sh


def hrule(slide, x, y, w, color=RULE, pt=0.75):
    """水平发丝线。pt → 英寸在这里转换，不把 Pt() 往外传。"""
    return _rect(slide, x, y, w, pt / 72.0, color)


def vrule(slide, x, y, h, color=RULE, pt=0.75):
    return _rect(slide, x, y, pt / 72.0, h, color)


def dot(slide, x, y, size=0.10, color=RED):
    return _rect(slide, x, y, size, size, color)


def outline_box(slide, x, y, w, h, color=DARK, radius=0.10):
    """1px 描边、无填充的容器（zcode 明确许可的形态）。"""
    return _rect(slide, x, y, w, h, color, outline=True, radius=radius)


def tint_band(slide, y, h, color=TINT):
    """通栏浅色带：左右出血到画布边缘，无圆角、无边框、无阴影。"""
    return _rect(slide, 0, y, CANVAS_W, h, color)


PALETTE = None  # 延迟填充，避免与上面的常量定义顺序耦合


def resolve_color(c):
    """颜色可以是 RGBColor，也可以是令牌名（'RED' / 'BLUE' / ...）。

    内容 spec 里写 `{'color': 'MUTED'}` 比写 RGBColor(0x64,0x64,0x63) 可读得多。
    """
    global PALETTE
    if PALETTE is None:
        PALETTE = dict(RED=RED, BLUE=BLUE, DARK=DARK, MUTED=MUTED,
                       GREY=GREY, RULE=RULE, TINT=TINT, WHITE=WHITE)
    if isinstance(c, str):
        try:
            return PALETTE[c.upper()]
        except KeyError:
            raise KeyError('unknown color name %r; known: %s'
                           % (c, sorted(PALETTE)))
    return c


def _as_text(v) -> str:
    """把「文本」字段统一成字符串。

    内容来自模型，字段类型不能假设：同一个字段可能是 `"文本"`、`["文本", {}]`、
    或 `[["文本", {}], ["续", {}]]`。递归拼平，避免 `r.text = list` 那种崩法。
    """
    if isinstance(v, str):
        return v
    if isinstance(v, (list, tuple)):
        return ''.join(_as_text(x) for x in v)
    return '' if v is None else str(v)


def put(slide, x, y, w, h, paras, *, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, ls=None, sa=None):
    """写文本。

    paras 结构：`[[(text, opts), ...], ...]`
      外层 = 段落，内层 = 段内的 run。
      opts 支持 size / bold / color（颜色对象或令牌名）/ ea / lat。
    """
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, 0)
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if ls is not None:
            p.line_spacing = ls
        if sa is not None:
            p.space_after = Pt(sa)
        for txt, o in para:
            r = p.add_run()
            r.text = _as_text(txt)
            r.font.size = Pt(o.get('size', FS['body']))
            r.font.bold = o.get('bold', False)
            r.font.color.rgb = resolve_color(o.get('color', DARK))
            r.font.name = o.get('lat', LAT)
            rPr = r._r.get_or_add_rPr()
            el = rPr.find(qn('a:ea'))
            if el is None:
                el = rPr.makeelement(qn('a:ea'), {})
                rPr.append(el)
            el.set('typeface', o.get('ea', EA))
    return tf


# ── 页面级组件 ────────────────────────────────────────────────

def header(slide, kicker=None, title=None):
    if kicker:
        put(slide, LEFT, Y_KICKER, W, 0.28,
            [[(kicker, dict(size=FS['kicker'], color=RED, bold=True))]])
    if title:
        put(slide, LEFT, Y_TITLE, W, 0.62,
            [[(title, dict(size=FS['title'], color=DARK, bold=True))]])


# 本次生成的截断记录。截断是**静默**的：它消除了溢出，于是几何检查全绿，
# 问题只在你肉眼看渲染图时才暴露（历史案例：已删除的 node_flow 8 个节点里 6 个
# 被截成 `小红书正文 · 爆款写作…`，而几何报告是干净的）。记下来，由 build()
# 收走写进日志。
TRUNCATIONS: list[tuple[str, str]] = []


def take_truncations() -> list[tuple[str, str]]:
    """取走并清空截断记录，返回 [(原文, 截断后), ...]。"""
    got = list(TRUNCATIONS)
    TRUNCATIONS.clear()
    return got


# ── 文字量宽（`_fit` 系列与 `wrap_lines` 共用同一个模型）─────────
# 中文/全角 1.0em，西文 0.52em。**这是估算**，所以调用方要留余量。
#
# 这几个码点在 CJK 字体里就是全角，但都小于 U+2E80，早先一律按 0.52 算 ——
# 而分隔页/封面的长标题里全是 `—` 和 `·`
# （`01 初识 Dify　—　什么是 Dify · 设计初衷 · 九大核心理念`），
# 于是宽度被低估 5–8%：两行折到贴边时，渲染出来就是溢出。
_WIDE_PUNCT = frozenset('—–…‘’“”·')


def _char_w(ch, pt):
    return pt * (1.0 if ch in _WIDE_PUNCT or ord(ch) > 0x2E80 else 0.52)


def text_w_in(text, pt):
    """估算一段文字占用的宽度（英寸）。"""
    return sum(_char_w(c, pt) for c in str(text)) / 72.0


def _cut(text, max_in, pt, lines=1):
    """裁到 `lines` 行内放得下，返回 (结果, 是否截断)。不记录。"""
    limit = max_in * 72.0 * max(lines, 1)
    w = 0.0
    for i, ch in enumerate(text):
        w += _char_w(ch, pt)
        if w > limit:
            return text[:max(i - 1, 1)].rstrip() + '…', True
    return text, False


def fit_one_line(text: str, max_in: float, size_pt: float, lines: int = 1) -> str:
    """把文本裁到 `lines` 行内放得下（超出加省略号）。

    有些框的高度只够固定行数（来源行 0.30" 只够 1 行；脚注 0.60" 够 2 行），
    模型一旦写长就会折行溢出。几何检查只能报「装不下」，修复回环又可能因为
    「不得改动数字」而改不动 —— 这类**确定性**超标由程序截断最可靠。
    """
    out, cut = _cut(str(text), max_in, size_pt, lines)
    if cut:
        TRUNCATIONS.append((str(text), out))
    return out


# ── 折行（标题这类**可以占两行**的框用它，而不是截断）───────────

# 宽度估算留的余量。折行是**硬换行**（写死段落），估偏了 PowerPoint 会再折一次，
# 于是「两行」变成三行、压到下面的内容 —— 宁可少放几个字。
_WIDTH_SAFETY = 0.95

# 可以断在**其后**的字符：空格、破折号、间隔号、各类标点。
_BREAK_AFTER = ' 　—–-·、，。；：！？,.;:!?)]}）】》」』%'
# 不能出现在**行首**的字符（中文避头规则）。找不到优先断点时按字断，
# 断点撞上这些字符就往回让 —— 否则行首会挂一个逗号。
_NO_LINE_START = '、，。；：！？）】》」』·—…,.;:!?)]}%'


def _last_break(cur: str) -> int:
    """`cur` 里最后一个优先断点（返回断点后的下标）；没有返回 -1。

    断点后面那个字不能是避头标点 —— `Dify · 设计初衷` 在空格后断，下一行就会
    以 `·` 开头；宁可退到前一个 `·` 之后断。
    """
    for j in range(len(cur) - 1, 0, -1):
        if cur[j - 1] in _BREAK_AFTER and cur[j] not in _NO_LINE_START:
            return j
    return -1


def _is_word_char(ch: str) -> bool:
    """拉丁字母/数字（含词内常见符号）—— 折行不能把它们劈开。"""
    return ord(ch) <= 0x2E80 and (ch.isalnum() or ch in '.-_+/#&@')


def _fallback_break(cur: str) -> int:
    """没有优先断点时按字断，但要躲开两个坑。

    ② 拉丁单词不许劈开（`Dify` 不能断成 `Dif` + `y`）
    ③ 行首不许挂避头标点（下一行不能以 `、` `，` `·` 开头）
    """
    n = len(cur)
    cut = n
    while cut > 1 and _is_word_char(cur[cut - 1]):
        cut -= 1
    if cut <= 1:
        cut = n
    while cut > 1 and cut < n and cur[cut] in _NO_LINE_START:
        cut -= 1
    return max(cut, 1)


def _wrap_at(text: str, w_in: float, pt: float) -> list[str]:
    """按给定行宽贪心折行。

    断点优先级：标点/空格之后 → 拉丁单词之前 → 任意字之间（避开行首标点）。
    """
    limit = w_in * 72.0
    lines, cur, cur_w, i = [], '', 0.0, 0
    while i < len(text):
        ch = text[i]
        cw = _char_w(ch, pt)
        if cur and cur_w + cw > limit:
            cut = _last_break(cur)
            if cut <= 0:
                cut = _fallback_break(cur)
            # 下一行的首字不能是避头标点 —— 标点宁可**吊在上一行末尾**
            # （中文排版的标点悬挂），也不能甩到下一行行首。
            while cut > 1:
                nxt = cur[cut] if cut < len(cur) else ch
                if nxt not in _NO_LINE_START:
                    break
                cut -= 1
            lines.append(cur[:cut].rstrip())
            cur = cur[cut:].lstrip(' ')    # 续行的行首空格不留（视觉缩进）
            cur_w = sum(_char_w(c, pt) for c in cur)
            continue                      # 同一个字符要重新试一次，别吞掉
        cur += ch
        cur_w += cw
        i += 1
    if cur.strip():
        lines.append(cur.rstrip())
    return lines or ['']


def wrap_lines(text: str, avail_in: float, pt: float, max_lines: int | None = None
               ) -> list[str]:
    """把文字折成若干行。

    `max_lines` 给了就**尽量把行拉匀**：先贪心折出最少行数，再二分找「还能装进
    `max_lines` 行的最小行宽」重新折 —— 否则 30 字标题会断成 29 字 + 1 字。
    返回的行数**不保证** ≤ `max_lines`（真装不下时由 `fit_block` 决定怎么办）。
    """
    text = str(text)
    limit = avail_in * _WIDTH_SAFETY
    lines = _wrap_at(text, limit, pt)
    if max_lines and len(lines) > max_lines:
        lo, hi = limit / float(max_lines), limit
        for _ in range(12):
            mid = (lo + hi) / 2.0
            if len(_wrap_at(text, mid, pt)) <= max_lines:
                hi = mid
            else:
                lo = mid
        lines = _wrap_at(text, hi, pt)
    return lines


def fit_block(text: str, avail_in: float, avail_h: float, *, sizes,
              max_lines: int = 2, ls: float = 1.2) -> tuple[list[str], float]:
    """在字号阶梯里挑最大的一个，让文字装进 `avail_in × avail_h` 且不超过 max_lines 行。

    返回 `(各行文字, 字号)`。**用折行换字号，而不是截断** —— 标题横排两行比
    「一行大字 + 省略号」好看得多，也不丢信息。

    阶梯全试完仍装不下（标题被改成极端长度）才截断：用最小字号折到 `max_lines` 行，
    末行加省略号，并把 `(原文, 截断后)` 记进 `TRUNCATIONS`。
    所以**阶梯的下限要低到实际到不了**，截断只该是理论上的一格保险。
    """
    text = str(text)
    sizes = list(sizes)
    for size in sizes:
        lines = wrap_lines(text, avail_in, size, max_lines=max_lines)
        if len(lines) <= max_lines and len(lines) * size * ls / 72.0 <= avail_h:
            return lines, size
    size = sizes[-1]
    lines = wrap_lines(text, avail_in, size, max_lines=max_lines)
    if len(lines) > max_lines:
        # 阶梯全试完还是装不下：只留前 max_lines 行，末行加省略号表示「后面还有」。
        # 末行本来就是放得下的，所以要**先**把省略号算进宽度再裁。
        keep = lines[:max_lines]
        tail, _ = _cut(keep[-1] + '…', avail_in, size)
        TRUNCATIONS.append((text, ''.join(keep[:-1]) + tail))
        lines = keep[:-1] + [tail]
    return lines, size


def footer(slide, page_no, source=None):
    if source:
        # 来源行框高仅 0.30"，必须单行放下 —— 超长直接截断
        src = fit_one_line(str(source), W - 1.4, FS['source'])
        put(slide, LEFT, Y_SOURCE, W - 1.4, 0.30,
            [[(src, dict(size=FS['source'], color=MUTED))]])
    # 页码框宽度取 1.10 而不是照抄模板占位符的 2.51 —— 模板那个右边界是 13.34"，
    # 已经越出 13.33" 画布，且右对齐后数字会落进装饰弧线区。
    put(slide, 10.83, Y_PAGENUM, 1.10, 0.28,
        [[(str(page_no), dict(size=FS['source'], color=MUTED))]],
        align=PP_ALIGN.RIGHT)

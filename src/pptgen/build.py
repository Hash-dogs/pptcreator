# -*- coding: utf-8 -*-
"""把内容 spec 组装成 pptx：填好公司模板的封面 / 目录，中间插入生成的正文页，封底保持原样。

⚠️ 核心工程约束（实测踩出来的）：

  **所有 `add_slide` 必须在删页之前完成。**

  python-pptx 的 `add_slide` 按「现有 slide 部件数量」推算新文件名。若先删掉模板里的
  内容页（部件数 4→3）再新增，新页会被命名成 `slide4.xml`，与既有封底**撞名** ——
  产出 zip 内出现两个同名条目、两个 rId 指向同一部件，PowerPoint 表现为「修复」或
  显示错页，而**模板封底会被静默覆盖丢失**。

  先加后删即可避免。`selfcheck()` 会把这条固化成可回归的断言。

模板三页的处理原则（第二版，实测后改的）：

  封面    填 deck 标题 + 日期。早先**一个字都不填** —— `spec['title']` 从
          `pipeline` 一路传到这里，却没人消费，成品第一页是张空白品牌图。
  目录    填真正的目录。**只写文字，不动格式** —— 模板自己的排版是深灰通栏上的
          40pt 白字 + 品牌红圆点列表，早先这里把标题拍成 28pt `DARK`，
          深色压深色几乎看不见；条目也被压到 16pt。
  封底    原样保留（品牌收尾页，没有可补的内容）。
"""
from __future__ import annotations
import collections
import datetime
import math
import os
import re
import zipfile

from pptx import Presentation
from pptx.util import Pt
from pptx.oxml.ns import qn

from .tokens import LAYOUT_BLANK, EA, EA_LIGHT, fit_block, take_truncations, text_w_in
from .layouts import render_slide


class BuildError(RuntimeError):
    pass


# 模板封面/目录页在 layout 里的名字 —— 用来**校验**我们填的是不是那两页。
# 靠 idx 定位占位符、靠 layout 名确认身份：模板换版时会在这里炸掉，
# 而不是把标题静默填到别的页上。
LAYOUT_COVER = 'Primary Title Slide'
LAYOUT_AGENDA = 'Agenda'

# 目录条目在模板里的字号（`slideLayout9.xml` 的 body lvl1 `sz=2400`）。
# python-pptx 读不到「继承来的字号」（它只解析本部件的 XML），而我们刻意不覆盖
# 字号，所以这里写死一个常量，专门用来算目录页**装不装得下**。
AGENDA_BODY_PT = 24
# 目录条目左边有 master bodyStyle 的 `marL=0.35"` 悬挂缩进（红点挂在这里），
# 真实可用宽要扣掉它。
AGENDA_INDENT_IN = 0.35


def _layout_by_name(prs, name):
    for l in prs.slide_layouts:
        if l.name == name:
            return l
    raise BuildError('layout %r not found in template; have: %s'
                     % (name, [l.name for l in prs.slide_layouts]))


def _ph(slide, idx):
    """按占位符 idx 取形状。封面/目录页靠 idx 定位，不靠 layout 名。"""
    for sh in slide.shapes:
        if sh.is_placeholder and sh.placeholder_format.idx == idx:
            return sh
    return None


def _check_layout(slide, want, what):
    got = slide.slide_layout.name
    if got != want:
        raise BuildError('模板的第 %s 页 layout 是 %r，期望 %r —— 模板换版了？'
                         % (what, got, want))


def _avail(shape):
    """占位符**真正可用**的宽高（英寸）。必须扣内边距，且只能量不能猜。"""
    emu = 914400.0
    tf = shape.text_frame
    w = shape.width / emu - (tf.margin_left + tf.margin_right) / emu
    h = shape.height / emu - (tf.margin_top + tf.margin_bottom) / emu
    return w, h


def _fill(shape, lines, *, size=None, ea=EA):
    """往占位符里写文字：**只设字号与中文字体，其余格式全部继承 layout**。

    字号以外的样式（颜色、对齐、加粗、行距、项目符号）模板里都已经调好了 ——
    早先 `fill_agenda` 手写 `28pt/DARK/bold`，把模板的白字标题改成深色压在深灰
    通栏上，几乎看不见。所以这里只做两件模板做不到的事：给字号（封面要按标题
    长度自适应）、补 `a:ea`（公司模板主题的 ea 为空，不补中文会走系统默认）。
    """
    tf = shape.text_frame
    tf.text = lines[0] if lines else ''
    for ln in lines[1:]:
        tf.add_paragraph().text = ln
    for p in tf.paragraphs:
        for r in p.runs:
            if size:
                r.font.size = Pt(size)
            rPr = r._r.get_or_add_rPr()
            el = rPr.find(qn('a:ea'))
            if el is None:
                el = rPr.makeelement(qn('a:ea'), {})
                rPr.append(el)
            el.set('typeface', ea)
    return tf


def fill_cover(prs, title, subtitle=''):
    """封面：deck 标题（过长折两行，不截断）+ 日期。

    模板封面的两个占位符（`ctrTitle` / `subTitle`）**一直是空的** —— 成品第一页
    只有 MEVION 的机器图，一个字都没有。`spec['title']` 从 `pipeline` 一路传到这里，
    却从来没人消费它。
    """
    title = (title or '').strip()
    if not title:
        # 没有标题就别动封面：只写一行日期看起来比空白更奇怪。
        # CLI 的 `build --content dify` 那条路径的 spec 就没有 title。
        return
    slide = prs.slides[0]
    _check_layout(slide, LAYOUT_COVER, '一页（封面）')
    ph = _ph(slide, 0)
    if ph is not None:
        # 模板自带 `normAutofit`，但它只在 PowerPoint 里编辑时才重算 ——
        # 而**封面在几何检查的 skip 名单里**（qa/geometry.py 默认跳过第 1/2/最后一页），
        # 溢出了没有任何东西报警。所以字号必须我们自己按可用宽算准。
        lines, size = fit_block(title, *_avail(ph),
                                sizes=(54, 48, 44, 40, 36, 32, 28, 24, 20))
        _fill(ph, lines, size=size, ea=EA_LIGHT)
    ph = _ph(slide, 1)
    if ph is not None and subtitle:
        _fill(ph, [subtitle])


def today_cn() -> str:
    """封面副标题的缺省值。"""
    d = datetime.date.today()
    return '%d 年 %d 月' % (d.year, d.month)


def fill_agenda(prs, items, on_log=None):
    """把目录页的占位符填成章节列表。items = [str, ...]

    标题固定写「目录」；字号/颜色/项目符号全部继承模板的 Agenda layout。
    `items` 为空也要调 —— 否则模板自带的「议题一/议题二/议题三」会留在成品里。
    """
    slide = prs.slides[1]
    _check_layout(slide, LAYOUT_AGENDA, '二页（目录）')
    title_ph = _ph(slide, 0)
    body_ph = next((sh for sh in slide.shapes
                    if sh.is_placeholder and sh.placeholder_format.idx != 0), None)
    if title_ph is not None:
        _fill(title_ph, ['目录'], ea=EA_LIGHT)
    if body_ph is None:
        return
    items = [str(t).strip() for t in (items or []) if str(t).strip()]
    if items and on_log:
        # 目录页**不在几何检查范围内**（geometry.analyse 默认 skip={1,2,n}），
        # 这里是唯一的溢出防线：按模板的 24pt 与 master 的 8pt 段前距估算高度。
        w, h = _avail(body_ph)
        w -= AGENDA_INDENT_IN
        rows = sum(max(1, math.ceil(text_w_in(t, AGENDA_BODY_PT) / max(w, 0.5)))
                   for t in items)
        need = (rows * AGENDA_BODY_PT + len(items) * 8) / 72.0
        if need > h:
            on_log('[build] 目录 %d 条估算需要 %.2f"，占位符只有 %.2f" —— '
                   '条目可能溢出（目录页不做几何检查）' % (len(items), need, h))
    _fill(body_ph, items)


def build(spec: dict, template: str, out_path: str, *, fill_toc: bool = True,
          on_log=None) -> str:
    """spec = {"slides": [...], "toc": [...], "title": str, "subtitle": str}"""
    # 截断记录是**模块级全局**：Web 端是 ThreadingHTTPServer + 长驻进程，
    # 上一轮（甚至并发另一个 job）的残留会串进这次的日志。开工先丢掉旧的。
    take_truncations()

    prs = Presentation(template)
    sldIdLst = prs.slides._sldIdLst
    n_orig = len(prs.slides)
    blank = _layout_by_name(prs, LAYOUT_BLANK)

    # ① 先全部新增（必须在删页之前！）
    #    页码自动分配：封面=1、目录=2，正文从 3 开始，封底不编号。
    for i, sl in enumerate(spec['slides']):
        sl['page'] = i + 3
        render_slide(prs.slides.add_slide(blank), sl)

    # ② 再删掉模板自带的空内容页（索引 2），保留封面 0 / 目录 1 / 封底 3
    if n_orig >= 4:
        ids = list(sldIdLst)
        dead = ids[2]
        prs.part.drop_rel(dead.get(qn('r:id')))
        sldIdLst.remove(dead)

    # ③ 重排：封面 / 目录 / 正文… / 封底
    ids = list(sldIdLst)
    order = [0, 1] + list(range(3, len(ids))) + [2]
    for e in list(sldIdLst):
        sldIdLst.remove(e)
    for i in order:
        sldIdLst.append(ids[i])

    fill_cover(prs, spec.get('title') or '', spec.get('subtitle') or today_cn())
    if fill_toc:
        # 目录为空也要调：否则模板自带的「议题一/议题二/议题三」会留在成品里。
        fill_agenda(prs, spec.get('toc') or [], on_log=on_log)

    # 截断是静默的（它消除了溢出，几何检查于是全绿），所以在这里把它捞出来。
    cuts = take_truncations()
    if cuts and on_log:
        on_log('[build] %d 处文案超长被截断：%s'
               % (len(cuts), '；'.join('%s → %s' % c for c in cuts[:4])))

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    prs.save(out_path)

    problems = selfcheck(out_path, expected_slides=len(ids))
    if problems:
        raise BuildError('结构自检未通过:\n  ' + '\n  '.join(problems))
    return out_path


def selfcheck(path: str, expected_slides: int | None = None) -> list[str]:
    """结构自检：zip 重名条目 / sldIdLst 目标唯一 / 图片 rel 可解析。"""
    problems: list[str] = []
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        dup = [n for n, c in collections.Counter(names).items() if c > 1]
        if dup:
            problems.append('zip 内存在重名条目（partname 撞名）: %s' % dup)

        pres = z.read('ppt/presentation.xml').decode('utf-8')
        rels = z.read('ppt/_rels/presentation.xml.rels').decode('utf-8')
        r2t = dict(re.findall(
            r'Id="(rId\d+)"[^>]*Target="slides/(slide\d+\.xml)"', rels))
        rids = re.findall(r'<p:sldId [^>]*r:id="(rId\d+)"', pres)
        targets = [r2t.get(r, '??') for r in rids]
        if len(set(targets)) != len(targets):
            problems.append('多个 sldId 指向同一个 slide 部件: %s' % targets)
        if expected_slides is not None and len(targets) != expected_slides:
            problems.append('页数不符: sldIdLst=%d, 期望=%d'
                            % (len(targets), expected_slides))
        for t in set(targets):
            if t != '??' and ('ppt/slides/' + t) not in names:
                problems.append('sldId 指向不存在的部件: %s' % t)

    # 图片 rel 断链（空框的成因）
    prs = Presentation(path)
    for i, slide in enumerate(prs.slides, 1):
        for rel in slide.part.rels.values():
            if rel.is_external:
                continue
            tp = getattr(rel, 'target_part', None)
            if tp is None:
                problems.append('s%d 存在无法解析的关系: %s' % (i, rel.rId))
            elif str(rel.reltype).endswith('/image') and not getattr(tp, 'blob', None):
                problems.append('s%d 图片关系 %s 的目标为空（会渲染成空框）' % (i, rel.rId))
    return problems
